"""Tests for hyperlink post-processing on Excel output."""

import openpyxl
import pandas as pd
import pytest

from tests.conftest import make_jpeg_with_exif

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _write_workbook(path, sheets):
    """Write one xlsx with the given {sheet_name: DataFrame} mapping."""
    with pd.ExcelWriter(path) as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)


def _sample_df(files=("a.jpg", "b.jpg")):
    return pd.DataFrame(
        {
            "Files": list(files),
            "Dates": ["01/01/2024"] * len(files),
            "Time": ["12:00:00"] * len(files),
        }
    )


# ---------------------------------------------------------------------------
# add_hyperlinks -- direct unit tests
# ---------------------------------------------------------------------------


def test_add_hyperlinks_sets_link_on_each_filename(extractor, tmp_path):
    directory = tmp_path / "photos"
    directory.mkdir()
    df = _sample_df()
    extractor.results.append((df, "Sheet", directory))
    output = tmp_path / "out.xlsx"
    _write_workbook(output, {"Sheet": df})

    extractor.add_hyperlinks(output.as_posix())

    wb = openpyxl.load_workbook(output)
    ws = wb["Sheet"]
    assert ws["A2"].hyperlink is not None
    assert ws["A2"].hyperlink.target == (directory / "a.jpg").as_posix()
    assert ws["A3"].hyperlink is not None
    assert ws["A3"].hyperlink.target == (directory / "b.jpg").as_posix()


def test_add_hyperlinks_skips_empty_cells(extractor, tmp_path):
    directory = tmp_path / "photos"
    directory.mkdir()
    df = _sample_df(files=("a.jpg", None, "c.jpg"))
    extractor.results.append((df, "Sheet", directory))
    output = tmp_path / "out.xlsx"
    _write_workbook(output, {"Sheet": df})

    extractor.add_hyperlinks(output.as_posix())

    wb = openpyxl.load_workbook(output)
    ws = wb["Sheet"]
    assert ws["A2"].hyperlink is not None
    assert ws["A3"].hyperlink is None  # blank cell skipped
    assert ws["A4"].hyperlink is not None


def test_add_hyperlinks_skips_none_directory(extractor, tmp_path):
    df = _sample_df(files=("a.jpg",))
    extractor.results.append((df, "Sheet", None))
    output = tmp_path / "out.xlsx"
    _write_workbook(output, {"Sheet": df})

    extractor.add_hyperlinks(output.as_posix())

    wb = openpyxl.load_workbook(output)
    assert wb["Sheet"]["A2"].hyperlink is None
    assert "Hyperlink Error" not in extractor.collected_warnings


def test_add_hyperlinks_handles_missing_sheet(extractor, tmp_path):
    directory = tmp_path / "photos"
    directory.mkdir()
    df = _sample_df(files=("a.jpg",))
    extractor.results.append((df, "Sheet", directory))
    output = tmp_path / "out.xlsx"
    _write_workbook(output, {"DifferentSheet": df})  # name mismatch

    extractor.add_hyperlinks(output.as_posix())

    assert "Hyperlink Error" in extractor.collected_warnings
    assert any("Sheet" in w for w in extractor.collected_warnings["Hyperlink Error"])


def test_add_hyperlinks_handles_missing_file(extractor, tmp_path):
    extractor.add_hyperlinks((tmp_path / "nope.xlsx").as_posix())
    assert "Hyperlink Error" in extractor.collected_warnings


def test_add_hyperlinks_applies_font_style(extractor, tmp_path):
    directory = tmp_path / "photos"
    directory.mkdir()
    df = _sample_df(files=("a.jpg",))
    extractor.results.append((df, "Sheet", directory))
    output = tmp_path / "out.xlsx"
    _write_workbook(output, {"Sheet": df})

    extractor.add_hyperlinks(output.as_posix())

    wb = openpyxl.load_workbook(output)
    cell = wb["Sheet"]["A2"]
    assert cell.font.color is not None
    assert cell.font.color.rgb == "000563C1"  # openpyxl stores as ARGB


# ---------------------------------------------------------------------------
# Integration: create_sheet honors hyperlink_var
# ---------------------------------------------------------------------------


@pytest.mark.real_create_sheet
def test_create_sheet_adds_hyperlinks_when_enabled(extractor, photos_only, tmp_path):
    extractor.hyperlink_var.set(True)
    extractor.keepopen_var.set(True)
    output = tmp_path / "output.xlsx"

    extractor.sort_and_extract(
        photos_only,
        recursive=False,
        final_output=output,
        modifier=0,
        output_series=True,
    )

    assert output.exists()
    wb = openpyxl.load_workbook(output)
    ws = wb["Sheet"]
    assert ws["A2"].hyperlink is not None
    assert ws["A2"].hyperlink.target == (photos_only / "IMG_001.jpg").as_posix()


@pytest.mark.real_create_sheet
def test_create_sheet_skips_hyperlinks_when_disabled(extractor, photos_only, tmp_path):
    extractor.hyperlink_var.set(False)
    extractor.keepopen_var.set(True)
    output = tmp_path / "output.xlsx"

    extractor.sort_and_extract(
        photos_only,
        recursive=False,
        final_output=output,
        modifier=0,
        output_series=True,
    )

    assert output.exists()
    wb = openpyxl.load_workbook(output)
    assert wb["Sheet"]["A2"].hyperlink is None


@pytest.mark.real_create_sheet
def test_recursive_mode_links_each_sheet_to_its_own_directory(extractor, tmp_path):
    """Regression: earlier loop nesting overwrote every sheet with the last dir."""
    extractor.hyperlink_var.set(True)
    extractor.keepopen_var.set(True)

    parent = tmp_path / "parent"
    for sub in ["Trip_A", "Trip_B"]:
        folder = parent / sub
        folder.mkdir(parents=True)
        make_jpeg_with_exif(folder / "IMG_001.jpg", "2024:05:15 12:30:45")

    output = tmp_path / "output.xlsx"
    extractor.maxdepth = 1
    extractor.find_directories(parent, output, modifier=0)
    assert output.exists()
    wb = openpyxl.load_workbook(output)

    ws_a = wb["Trip_A"]
    assert ws_a["A2"].hyperlink.target == (parent / "Trip_A" / "IMG_001.jpg").as_posix()

    ws_b = wb["Trip_B"]
    assert ws_b["A2"].hyperlink.target == (parent / "Trip_B" / "IMG_001.jpg").as_posix()
